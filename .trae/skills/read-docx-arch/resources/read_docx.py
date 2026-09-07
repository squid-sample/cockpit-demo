#!/usr/bin/env python3
"""读取 .docx 文档的文字、表格和嵌入图片，输出结构化结果。"""

import json
import os
import sys
from pathlib import Path


def extract_docx(file_path: str):
    try:
        from docx import Document
    except ImportError:
        print(json.dumps({
            "success": False,
            "error": "缺少 python-docx，请运行: pip install python-docx"
        }, ensure_ascii=False))
        sys.exit(1)

    p = Path(file_path)
    if not p.exists():
        print(json.dumps({
            "success": False,
            "error": f"文件不存在: {file_path}"
        }, ensure_ascii=False))
        sys.exit(1)

    doc = Document(file_path)
    base_dir = p.parent
    images_dir = base_dir / f"{p.stem}.docx-images"
    images_dir.mkdir(exist_ok=True)

    # 1. 提取文字（保留标题层级）
    paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style_name = para.style.name if para.style else ""
        if style_name.startswith("Heading"):
            level = style_name.replace("Heading", "").strip()
            prefix = "#" * (int(level) if level.isdigit() else 1)
            paragraphs.append(f"{prefix} {text}")
        else:
            paragraphs.append(text)

    # 2. 提取表格
    tables = []
    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(" | ".join(cells))
        tables.append("\n".join(rows))

    # 3. 提取图片
    image_files = []
    image_index = 0
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            image_index += 1
            image = rel.target_part
            content_type = image.content_type
            ext = content_type.split("/")[-1]
            if ext == "jpeg":
                ext = "jpg"
            img_name = f"image_{image_index:02d}.{ext}"
            img_path = images_dir / img_name
            with open(img_path, "wb") as f:
                f.write(image.blob)
            image_files.append(str(img_path))

    result = {
        "success": True,
        "file_path": file_path,
        "title": p.stem,
        "paragraphs": paragraphs,
        "tables": tables,
        "image_count": len(image_files),
        "images_dir": str(images_dir),
        "image_files": image_files,
        "text_preview": "\n\n".join(paragraphs[:50])
    }

    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(json.dumps({
            "success": False,
            "error": "用法: python read_docx.py <docx文件路径>"
        }, ensure_ascii=False))
        sys.exit(1)
    extract_docx(sys.argv[1])
